import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, fill_hex):
    """Sets background color of a table cell."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    """Sets padding inside a table cell in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    """Applies subtle border lines."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def format_row(row, bg_hex, is_header=False):
    for cell in row.cells:
        set_cell_background(cell, bg_hex)
        set_cell_margins(cell, top=120 if is_header else 90, bottom=120 if is_header else 90)

def generate_full_report(output_path: str):
    doc = Document()

    # Configure Margins
    for sec in doc.sections:
        sec.top_margin = Inches(0.75)
        sec.bottom_margin = Inches(0.75)
        sec.left_margin = Inches(0.8)
        sec.right_margin = Inches(0.8)

        # Header
        h_p = sec.header.paragraphs[0]
        h_p.text = "MRPL OmniAI EngineCore — Audit & Compliance Intelligence Report"
        h_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if h_p.runs:
            h_p.runs[0].font.size = Pt(8.5)
            h_p.runs[0].font.color.rgb = RGBColor(148, 163, 184)

        # Footer
        f_p = sec.footer.paragraphs[0]
        f_p.text = "Mangalore Refinery and Petrochemicals Limited (MRPL) — Plant Maintenance SOP Audit"
        f_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if f_p.runs:
            f_p.runs[0].font.size = Pt(8.5)
            f_p.runs[0].font.italic = True
            f_p.runs[0].font.color.rgb = RGBColor(100, 116, 139)

    # Document Header
    p_title = doc.add_paragraph()
    r_title = p_title.add_run("PUMP INSPECTION & SOP COMPLIANCE AUDIT REPORT")
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 32, 67) # MRPL Deep Navy
    p_title.paragraph_format.space_after = Pt(2)

    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run("Automated Vision OCR & Document Intelligence Audit — Standard: SOP-MNT-042")
    r_sub.font.size = Pt(10.5)
    r_sub.font.color.rgb = RGBColor(71, 85, 105)
    p_sub.paragraph_format.space_after = Pt(12)

    # Status Banner Box
    banner_tbl = doc.add_table(rows=1, cols=1)
    banner_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    b_cell = banner_tbl.cell(0, 0)
    b_cell.width = Inches(6.9)
    set_cell_background(b_cell, "FEE2E2") # Red-100
    set_cell_margins(b_cell, top=140, bottom=140, left=180, right=180)
    
    bp = b_cell.paragraphs[0]
    r_b1 = bp.add_run("OVERALL AUDIT RESULT: FAILED — NON-COMPLIANT (DISPOSITION: REJECT)\n")
    r_b1.font.bold = True
    r_b1.font.size = Pt(12)
    r_b1.font.color.rgb = RGBColor(185, 28, 28)

    r_b2 = bp.add_run(
        "Audit Finding: Pump P-302B is operating in Zone C (Unsatisfactory, 6.8 mm/s RMS peak). "
        "Seal temperature of 84.6 °C exceeds the mandatory 80.0 °C threshold. Critical approval note requirements "
        "have failed: oil cleanliness rating missing, technician and certified supervisor signatures absent, "
        "and no maintenance work order created. Do not approve until all mandatory corrective actions are closed."
    )
    r_b2.font.size = Pt(9.5)
    r_b2.font.color.rgb = RGBColor(127, 29, 29)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 1. Executive Metadata Summary
    h1 = doc.add_heading("1. Inspection & Equipment Identification", level=2)
    h1.paragraph_format.space_before = Pt(10)
    h1.paragraph_format.space_after = Pt(4)

    meta_tbl = doc.add_table(rows=7, cols=2)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(meta_tbl)

    metadata_items = [
        ("Report Reference ID", "SYN-PIR-042-009"),
        ("Equipment Tag ID", "P-302B"),
        ("Plant Location / Area", "Process Pump House — Train 2, Bay B"),
        ("Inspection Date & Type", "02 September 2026 | Routine Vibration Inspection"),
        ("Inspector", "K. Arun — Synthetic Maintenance Technician"),
        ("Governing SOP Standard", "MRPL SOP-MNT-042 (Referenced ambiguously as 'MNT-042' in submission)"),
        ("OCR & Trace Origin", "MRPL OmniAI EngineCore (Trace #559447e8 | #826b783b)")
    ]

    for idx, (lbl, val) in enumerate(metadata_items):
        row = meta_tbl.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.4)
        c1.width = Inches(4.5)
        set_cell_background(c0, "F8FAFC")
        set_cell_margins(c0, 80, 80, 100, 100)
        set_cell_margins(c1, 80, 80, 100, 100)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(lbl)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = RGBColor(30, 41, 59)

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 2. Vibration Measurement Summary
    h2 = doc.add_heading("2. Vibration Measurement & Zone Classification", level=2)
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(4)

    vib_tbl = doc.add_table(rows=5, cols=5)
    vib_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(vib_tbl)

    vib_headers = ["Measurement Point", "Peak RMS (mm/s)", "Peak Freq (Hz)", "SOP Zone", "Required Response"]
    for j, h in enumerate(vib_headers):
        cell = vib_tbl.cell(0, j)
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, 100, 100, 80, 80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in [1, 2, 3] else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    vib_data = [
        ("DE Pump — Horizontal", "6.8", "117.4", "Zone C", "Inspection approval note + WO + lubrication check"),
        ("DE Pump — Vertical", "6.2", "116.9", "Zone C", "Inspection approval note + WO + lubrication check"),
        ("NDE Pump — Horizontal", "5.7", "118.1", "Zone C", "Inspection approval note + WO + lubrication check"),
        ("NDE Pump — Vertical", "4.9", "117.8", "Zone C", "Inspection approval note + WO + lubrication check")
    ]

    for i, row_data in enumerate(vib_data, start=1):
        row = vib_tbl.rows[i]
        bg = "FFFFFF" if i % 2 != 0 else "F8FAFC"
        for j, val in enumerate(row_data):
            cell = row.cells[j]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 80, 80, 80, 80)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in [1, 2, 3] else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            r.font.size = Pt(9)
            if j == 1 and float(val) >= 6.8:
                r.font.bold = True
                r.font.color.rgb = RGBColor(185, 28, 28)
            elif j == 3:
                r.font.bold = True
                r.font.color.rgb = RGBColor(234, 88, 12)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 3. SOP Threshold Criteria Evaluation
    h3 = doc.add_heading("3. SOP Threshold Criteria Evaluation (ISO 10816 / SOP-MNT-042)", level=2)
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(4)

    sop_tbl = doc.add_table(rows=5, cols=4)
    sop_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(sop_tbl)

    sop_headers = ["SOP Zone", "Criterion Threshold", "Result for P-302B", "Evaluation Status"]
    for j, h in enumerate(sop_headers):
        cell = sop_tbl.cell(0, j)
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, 100, 100, 80, 80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in [2, 3] else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    sop_data = [
        ("Zone A — Normal", "< 2.8 mm/s", "6.8 mm/s", "FAIL"),
        ("Zone B — Acceptable / Alert", "2.8 to 4.5 mm/s", "6.8 mm/s", "FAIL"),
        ("Zone C — Unsatisfactory", "4.5 to 7.1 mm/s", "6.8 mm/s", "TRIGGERED"),
        ("Zone D — Immediate Shutdown", "> 7.1 mm/s", "6.8 mm/s", "Not Triggered")
    ]

    for i, row_data in enumerate(sop_data, start=1):
        row = sop_tbl.rows[i]
        bg = "FFFFFF" if i % 2 != 0 else "F8FAFC"
        for j, val in enumerate(row_data):
            cell = row.cells[j]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 80, 80, 80, 80)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in [2, 3] else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            r.font.size = Pt(9)
            if j == 3:
                r.font.bold = True
                if val == "TRIGGERED":
                    r.font.color.rgb = RGBColor(234, 88, 12)
                elif val == "FAIL":
                    r.font.color.rgb = RGBColor(185, 28, 28)
                else:
                    r.font.color.rgb = RGBColor(22, 163, 74)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 4. Mandatory Approval Note Audit Check
    h4 = doc.add_heading("4. Mandatory Approval Note Parameter Check", level=2)
    h4.paragraph_format.space_before = Pt(10)
    h4.paragraph_format.space_after = Pt(4)

    appr_tbl = doc.add_table(rows=10, cols=3)
    appr_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(appr_tbl)

    appr_headers = ["Mandatory Requirement", "Submitted Record Entry", "Audit Assessment"]
    for j, h in enumerate(appr_headers):
        cell = appr_tbl.cell(0, j)
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, 100, 100, 80, 80)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    appr_data = [
        ("SOP-MNT-042 Explicit Reference", "Only listed as 'MNT-042' in remarks", "FAIL — Exact standard reference missing"),
        ("Equipment Tag ID", "P-302B", "PASS — Tag verified"),
        ("Location", "Process Pump House — Train 2, Bay B", "PASS — Location verified"),
        ("Peak Vibration RMS", "6.8 mm/s", "PASS (Recorded, exceeds Zone B)"),
        ("Peak Frequency", "117.4 Hz", "PASS (Dominant harmonic recorded)"),
        ("Oil Cleanliness Rating", "Not recorded", "FAIL — Mandatory rating absent"),
        ("Seal Temperature", "84.6 °C (> 80 °C threshold)", "FAIL — Exceeds 80.0 °C upper operating limit"),
        ("Maintenance Technician Signature", "Blank", "FAIL — Mandatory sign-off missing"),
        ("Certified Supervisor Endorsement", "Blank", "FAIL — Mandatory endorsement missing")
    ]

    for i, row_data in enumerate(appr_data, start=1):
        row = appr_tbl.rows[i]
        bg = "FFFFFF" if i % 2 != 0 else "F8FAFC"
        for j, val in enumerate(row_data):
            cell = row.cells[j]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 80, 80, 80, 80)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9)
            if j == 2:
                r.font.bold = True
                if val.startswith("FAIL"):
                    r.font.color.rgb = RGBColor(185, 28, 28)
                else:
                    r.font.color.rgb = RGBColor(22, 163, 74)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 5. Maintenance Action Tracking & Audit Violations
    h5 = doc.add_heading("5. Maintenance Action Tracking & SOP Violations", level=2)
    h5.paragraph_format.space_before = Pt(10)
    h5.paragraph_format.space_after = Pt(4)

    act_tbl = doc.add_table(rows=6, cols=4)
    act_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(act_tbl)

    act_headers = ["Action Required by SOP", "Regulatory Mandate", "Submitted Record", "Audit Status"]
    for j, h in enumerate(act_headers):
        cell = act_tbl.cell(0, j)
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, 100, 100, 80, 80)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    act_data = [
        ("Create Maintenance Work Order", "Mandatory for Zone C condition", "No work-order number entered", "FAIL"),
        ("Bearing Lubrication Check", "Mandatory for Zone C condition", "Not performed / not documented", "FAIL"),
        ("Vibration Re-Measurement", "Required corrective follow-up", "No follow-up date recorded", "FAIL"),
        ("Seal-Temperature Investigation", "Mandatory when seal temp > 80 °C", "No investigation recorded", "FAIL"),
        ("Technician & Supervisor Signatures", "Mandatory approval requirement", "Absent / Blank", "FAIL")
    ]

    for i, row_data in enumerate(act_data, start=1):
        row = act_tbl.rows[i]
        bg = "FFFFFF" if i % 2 != 0 else "F8FAFC"
        for j, val in enumerate(row_data):
            cell = row.cells[j]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 80, 80, 80, 80)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9)
            if j == 3:
                r.font.bold = True
                r.font.color.rgb = RGBColor(185, 28, 28)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 6. Corrective Action Plan (CAP) Matrix
    h6 = doc.add_heading("6. Proposed Corrective Action Plan (CAP Matrix)", level=2)
    h6.paragraph_format.space_before = Pt(10)
    h6.paragraph_format.space_after = Pt(4)

    cap_tbl = doc.add_table(rows=7, cols=4)
    cap_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(cap_tbl)

    cap_headers = ["CA ID", "Corrective Action Description", "Priority", "Mandatory Closure Evidence"]
    for j, h in enumerate(cap_headers):
        cell = cap_tbl.cell(0, j)
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, 100, 100, 80, 80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in [0, 2] else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    cap_data = [
        ("CA-042-01", "Create maintenance work order for P-302B vibration.", "High", "Valid Work Order ID in SAP/Maximo"),
        ("CA-042-02", "Perform bearing lubrication inspection and oil sampling.", "High", "Technician Log & Lube Quality Sheet"),
        ("CA-042-03", "Investigate mechanical seal temperature; restore below 80 °C.", "High", "Thermal Imaging & Re-test Report"),
        ("CA-042-04", "Perform and document oil cleanliness rating (ISO 4406).", "Medium", "Lab Oil Analysis Certificate"),
        ("CA-042-05", "Repeat multi-plane vibration re-measurement across DE/NDE.", "High", "Post-maintenance Spectrum Report"),
        ("CA-042-06", "Obtain signed approval note with certified supervisor endorsement.", "High", "Signed & Endorsed Inspection Form")
    ]

    for i, row_data in enumerate(cap_data, start=1):
        row = cap_tbl.rows[i]
        bg = "FFFFFF" if i % 2 != 0 else "F8FAFC"
        for j, val in enumerate(row_data):
            cell = row.cells[j]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 80, 80, 80, 80)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in [0, 2] else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            r.font.size = Pt(9)
            if j == 2 and val == "High":
                r.font.bold = True
                r.font.color.rgb = RGBColor(185, 28, 28)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 7. Final Assessment & Audit Disposition
    h7 = doc.add_heading("7. Final Audit Disposition & Sign-Off Requirement", level=2)
    h7.paragraph_format.space_before = Pt(10)
    h7.paragraph_format.space_after = Pt(4)

    p_disp = doc.add_paragraph()
    r_d1 = p_disp.add_run("FINAL DISPOSITION: REJECT / NON-COMPLIANT\n")
    r_d1.font.bold = True
    r_d1.font.color.rgb = RGBColor(185, 28, 28)
    r_d2 = p_disp.add_run(
        "Under MRPL Standard Operating Procedure SOP-MNT-042, this inspection cannot be approved. "
        "The equipment remains flagged for maintenance action. All six Corrective Actions (CA-042-01 to CA-042-06) "
        "must be executed, re-tested, and endorsed by an authorized plant supervisor before normal operating clearance is restored."
    )
    r_d2.font.size = Pt(9.5)
    r_d2.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Sign-off boxes table
    sign_tbl = doc.add_table(rows=2, cols=2)
    sign_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(sign_tbl)

    sign_cells = [
        ("Maintenance Technician Sign-off:\n\n___________________________\nName: K. Arun\nDate: ____/____/________",
         "Certified Area Supervisor Endorsement:\n\n___________________________\nName: ______________________\nDate: ____/____/________"),
        ("Compliance Auditor Signature:\n\nMRPL OmniAI Automated Engine\nStatus: NON-COMPLIANT AUDIT GENERATED",
         "Plant Reliability Department:\n\n___________________________\nAction: WORK ORDER PENDING")
    ]

    for row_i, (t1, t2) in enumerate(sign_cells):
        row = sign_tbl.rows[row_i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(3.45)
        c1.width = Inches(3.45)
        set_cell_background(c0, "F8FAFC")
        set_cell_background(c1, "F8FAFC")
        set_cell_margins(c0, 100, 100, 100, 100)
        set_cell_margins(c1, 100, 100, 100, 100)
        c0.paragraphs[0].text = t1
        c1.paragraphs[0].text = t2
        for r in c0.paragraphs[0].runs:
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(71, 85, 105)
        for r in c1.paragraphs[0].runs:
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(71, 85, 105)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Full comprehensive document generated successfully at: {output_path}")

if __name__ == "__main__":
    out_file = r"c:\sih117\prototype\MRPL_SOP_MNT_042_Pump_Inspection_Report.docx"
    generate_full_report(out_file)
